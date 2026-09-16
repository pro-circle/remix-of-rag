"""Optional: turn the sample .txt files into .docx / .pdf so those parsers can be demoed.

    python scripts/make_sample_binaries.py
"""
from __future__ import annotations

from pathlib import Path

SAMPLES = Path(__file__).resolve().parents[1] / "sample_data"


def to_docx(src: Path) -> None:
    import docx

    doc = docx.Document()
    for para in src.read_text().split("\n\n"):
        text = " ".join(para.split())
        if not text:
            continue
        first = text.split(" ")[0]
        if first[:1].isdigit() and first.rstrip(".").replace(".", "").isdigit():
            doc.add_heading(text[:100], level=min(4, first.count(".") + 1))
        else:
            doc.add_paragraph(text)
    doc.save(str(src.with_suffix(".docx")))


def to_pdf(src: Path) -> None:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(src.with_suffix(".pdf")), pagesize=LETTER)
    width, height = LETTER
    y = height - 60
    for line in src.read_text().split("\n"):
        if y < 60:
            c.showPage()
            y = height - 60
        c.setFont("Helvetica", 10)
        c.drawString(50, y, line[:105])
        y -= 14
    c.save()


if __name__ == "__main__":
    for path in sorted(SAMPLES.glob("*.txt")):
        try:
            to_docx(path)
            print("docx:", path.with_suffix(".docx").name)
        except ImportError:
            print("skip docx (pip install python-docx)")
        try:
            to_pdf(path)
            print("pdf :", path.with_suffix(".pdf").name)
        except ImportError:
            print("skip pdf (pip install reportlab)")
