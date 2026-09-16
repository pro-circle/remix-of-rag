from __future__ import annotations

from pathlib import Path

from ..utils.errors import RagError
from .base import Block, DocumentPage, ParsedDocument


class DocxParser:
    extensions = (".docx",)

    def parse(self, path: Path, name: str) -> ParsedDocument:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover
            raise RagError("DOCX support is not installed on the server.", 500, detail=str(exc))
        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise RagError("This Word file could not be read. It may be corrupted.", 422, detail=str(exc))

        blocks: list[Block] = []
        for para in document.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            style = (para.style.name or "").lower()
            if style.startswith("heading"):
                level = int(style.replace("heading", "").strip() or 1)
                blocks.append(Block(text=text, block_type="heading", level=min(4, level)))
            elif style.startswith("list") or style.startswith("bullet"):
                blocks.append(Block(text=f"- {text}", block_type="list"))
            elif style == "title":
                blocks.append(Block(text=text, block_type="heading", level=1))
            else:
                blocks.append(Block(text=text))

        for table in document.tables:
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
            rows = [r for r in rows if r.strip(" |")]
            if rows:
                blocks.append(Block(text="\n".join(rows), block_type="table"))

        text = "\n\n".join(b.text for b in blocks)
        page = DocumentPage(page_number=1, text=text, blocks=blocks)
        return ParsedDocument(name=name, pages=[page])
